#!/usr/bin/env python3
"""Robô do Habilita: lê a pasta pública do SerproDrive, baixa só o que é novo ou mudou,
extrai o texto (com leitura de imagem para PDF escaneado), descobre a validade de cada
certidão e gera o índice estático do site (docs/dados.json), com cópia dos PDFs e miniaturas.

Não usa nenhuma API paga: WebDAV público do Nextcloud + pypdf/pdfium + tesseract (se instalado).
Uso: python3 sync/sincronizar.py
"""
import base64, datetime as dt, hashlib, io, json, os, re, shutil, subprocess, sys, unicodedata
import urllib.parse, urllib.request, xml.etree.ElementTree as ET
from pathlib import Path

TOKEN = os.environ.get("SERPRODRIVE_TOKEN", "rYe4xqS7Tr3xwkS")
BASE = "https://serprodrive.serpro.gov.br"
RAIZ = Path(__file__).resolve().parent.parent
SITE = RAIZ / "docs"
CACHE = RAIZ / ".cache" / "arquivos"
DADOS = SITE / "dados.json"
HOJE = dt.date.today()

# ---------------------------------------------------------------- SerproDrive (WebDAV público)
AUTH = "Basic " + base64.b64encode(f"{TOKEN}:".encode()).decode()
CABECALHOS = {"Authorization": AUTH, "X-Requested-With": "XMLHttpRequest"}  # sem este cabeçalho o Nextcloud 20 nega
NS = {"d": "DAV:"}
PROPFIND = b'<?xml version="1.0"?><d:propfind xmlns:d="DAV:"><d:prop><d:getlastmodified/><d:getcontentlength/><d:resourcetype/><d:getetag/></d:prop></d:propfind>'


def webdav(caminho):
    return f"{BASE}/public.php/webdav" + urllib.parse.quote(caminho)


def listar(caminho="/"):
    req = urllib.request.Request(webdav(caminho), data=PROPFIND, method="PROPFIND",
                                 headers={**CABECALHOS, "Depth": "1", "Content-Type": "application/xml"})
    xml = ET.fromstring(urllib.request.urlopen(req, timeout=60).read())
    itens = []
    for r in xml.findall("d:response", NS):
        href = urllib.parse.unquote(r.find("d:href", NS).text).split("/public.php/webdav", 1)[1]
        if href.rstrip("/") == caminho.rstrip("/"):
            continue
        p = r.find("d:propstat/d:prop", NS)
        itens.append({
            "caminho": href,
            "pasta": p.find("d:resourcetype/d:collection", NS) is not None,
            "tamanho": int(p.findtext("d:getcontentlength", namespaces=NS) or 0),
            "modificadoEm": dt.datetime.strptime(p.findtext("d:getlastmodified", namespaces=NS), "%a, %d %b %Y %H:%M:%S %Z").isoformat() + "Z",
            "etag": (p.findtext("d:getetag", namespaces=NS) or "").strip('"'),
        })
    return itens


def arvore():
    todos, fila = [], ["/"]
    while fila:
        for item in listar(fila.pop()):
            todos.append(item)
            if item["pasta"]:
                fila.append(item["caminho"])
    return todos


def baixar(caminho, destino):
    destino.parent.mkdir(parents=True, exist_ok=True)
    req = urllib.request.Request(webdav(caminho), headers=CABECALHOS)
    destino.write_bytes(urllib.request.urlopen(req, timeout=180).read())


def link_download(caminho):
    pasta, nome = caminho.rsplit("/", 1)
    return f"{BASE}/s/{TOKEN}/download?path={urllib.parse.quote(pasta or '/')}&files={urllib.parse.quote(nome)}"


def link_pasta(pasta):
    return f"{BASE}/s/{TOKEN}?path={urllib.parse.quote(pasta or '/')}"


# ---------------------------------------------------------------- texto
def sem_acento(s):
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")


def texto_pdf(arquivo):
    """Texto das primeiras páginas; PDF escaneado (sem camada de texto) vai para a leitura de imagem."""
    from pypdf import PdfReader
    try:
        paginas = PdfReader(str(arquivo)).pages
        texto = "\n".join((p.extract_text() or "") for p in paginas[:8])
    except Exception:
        texto = ""
    if len(texto.strip()) >= 80:
        return texto, "texto"
    ocr = ocr_pdf(arquivo)
    return (ocr, "ocr") if ocr else (texto, "sem texto")


def ocr_pdf(arquivo):
    if not shutil.which("tesseract"):
        return ""
    import pypdfium2 as pdfium
    partes = []
    try:
        pdf = pdfium.PdfDocument(str(arquivo))
        for i in range(min(3, len(pdf))):
            img = pdf[i].render(scale=2.2).to_pil()
            buf = io.BytesIO(); img.save(buf, "PNG")
            r = subprocess.run(["tesseract", "stdin", "stdout", "-l", "por"], input=buf.getvalue(), capture_output=True, timeout=180)
            partes.append(r.stdout.decode("utf-8", "ignore"))
    except Exception as e:
        print("   OCR falhou:", arquivo.name, e)
    return "\n".join(partes)


def miniatura(arquivo, destino):
    import pypdfium2 as pdfium
    try:
        pdf = pdfium.PdfDocument(str(arquivo))
        img = pdf[0].render(scale=0.9).to_pil()
        img.thumbnail((360, 480))
        destino.parent.mkdir(parents=True, exist_ok=True)
        img.save(destino, "WEBP", quality=72)
        return True
    except Exception as e:
        print("   miniatura falhou:", arquivo.name, e)
        return False


# ---------------------------------------------------------------- validade
DATA = r"(\d{1,2}/\d{1,2}/\d{4})"


def para_data(s):
    try:
        d, m, a = (int(x) for x in s.split("/"))
        return dt.date(a, m, d)
    except Exception:
        return None


def validade_no_texto(texto, categoria):
    """Ordem pensada nos modelos reais da pasta (SICAF, RFB/PGFN, estaduais, municipais, FGTS, TST, TJDFT)."""
    t = re.sub(r"\s+", " ", sem_acento(texto).lower())
    regras = [
        rf"data de vencimento do cadastro:?\s*{DATA}",            # SICAF (o documento lista várias certidões; vale o cadastro)
        rf"(?:sera |e )?valida ate:?\s*(?:o dia\s*)?{DATA}",       # "Válida até 12/04/2027"
        rf"validade:?\s*{DATA}\s*(?:a|ate)\s*{DATA}",              # FGTS: "Validade: 18/10/2026 a 16/11/2026"
        rf"(?:data de )?validade:?\s*{DATA}",                      # TST: "Validade: 14/04/2027"
        rf"vencimento:?\s*{DATA}",
    ]
    if "sicaf" not in categoria.lower():
        regras = regras[1:]
    for regra in regras:
        m = re.search(regra, t)
        if m:
            data = para_data(m.groups()[-1])
            if data and 2000 < data.year < 2100:
                return {"ate": data.isoformat(), "fonte": "documento"}
    # "válida por 60 dias" a partir da emissão
    dias = re.search(r"valida por (\d{1,3})\s*(?:\([a-z ]+\)\s*)?dias", t)
    emissao = re.search(rf"(?:emitid[ao][^0-9]{{0,40}}|emissao:?\s*|data de emissao:?\s*){DATA}", t)
    if dias and emissao:
        base = para_data(emissao.group(1))
        if base:
            return {"ate": (base + dt.timedelta(days=int(dias.group(1)))).isoformat(), "fonte": "documento", "emissao": base.isoformat()}
    return None


def tabela_vencimentos(arquivo):
    """Planilha de controle: nome (cidade/UF ou estado) -> validade. Contatos internos ficam de fora."""
    import openpyxl
    mapa = {}
    try:
        wb = openpyxl.load_workbook(arquivo, data_only=True)
        aba = next((w for w in wb.worksheets if "complet" in w.title.lower()), wb.worksheets[0])
        secao = ""
        for linha in aba.iter_rows(values_only=True):
            nome, validade = (linha[0], linha[1]) if len(linha) > 1 else (None, None)
            if isinstance(nome, str) and nome.strip().upper().startswith("CERTID"):
                secao = "municipal" if "MUNIC" in nome.upper() else "estadual" if "ESTAD" in nome.upper() else nome.strip().lower()
                continue
            if isinstance(nome, str) and isinstance(validade, dt.datetime):
                mapa[(secao, chave(nome))] = validade.date().isoformat()
    except Exception as e:
        print("   planilha de vencimentos não lida:", e)
    return mapa


def chave(s):
    s = sem_acento(str(s)).upper()
    s = re.sub(r"\b(ESCRITORIO|REGIONAL|DE|DO|DA|CERTIDAO|ESTADUAL|MUNICIPAL)\b", " ", s)
    return re.sub(r"[^A-Z]", "", s.split("/")[0])


# ---------------------------------------------------------------- organização (categoria, local, título)
CATEGORIAS = {
    "SICAF": ("SICAF", "🏛️", "Declaração do Sistema de Cadastramento Unificado de Fornecedores, por regional/escritório."),
    "CERTIDAO DE REGULARIDADE FEDERAL": ("Certidão federal", "🧾", "Certidão conjunta de débitos federais (Receita Federal, INSS e PGFN)."),
    "CERTIDAO DE REGULARIDADE ESTADUAL": ("Certidão estadual", "🗺️", "Certidões de regularidade fiscal estadual, por UF."),
    "CERTIDAO DE REGULARIDADE MUNICIPAL": ("Certidão municipal", "🏙️", "Certidões de regularidade fiscal municipal, por cidade da regional/escritório."),
    "CERTIDAO DE REGULARIDADE FGTS": ("FGTS (CRF)", "🏦", "Certificado de Regularidade do FGTS emitido pela Caixa."),
    "CERTIDAO TRABALHISTA": ("Trabalhista (CNDT)", "⚖️", "Certidão Negativa de Débitos Trabalhistas do TST."),
    "CERTIDAO FALENCIA E RECUPERACAO JUDICIAL": ("Falência e recuperação judicial", "📜", "Certidão de falências e recuperações judiciais do TJDFT."),
    "CADIN": ("CADIN", "📋", "Situação no Cadastro Informativo de créditos não quitados do setor público federal."),
    "ESTATUTO SOCIAL DO SERPRO": ("Estatuto social", "📘", "Estatuto Social do Serpro."),
    "LEIS DE CRIACAO DO SERPRO": ("Leis de criação", "📖", "Leis de criação do Serpro (Lei nº 4.516/1964 e nº 5.615/1970) e base jurídica."),
    "DRE BALANCOS E CAPITAL SOCIAL": ("Balanços e DRE", "📊", "Balanço patrimonial, DRE e capital social."),
    "INFORMACOES SOBRE FATURAMENTO": ("Faturamento", "💰", "Informações sobre faturamento."),
    "TCU": ("TCU", "🔎", "Consultas e certidões do TCU (licitantes inidôneos, consulta consolidada)."),
    "IMUNIDADE": ("Imunidade tributária", "🛡️", "Atos declaratórios de imunidade tributária."),
    "CERTIDOES JUNTA COMERCIAL": ("Junta Comercial", "🏢", "Certidões da Junta Comercial do Distrito Federal."),
    "CNPJ SEDE REGIONAIS E ESCRITORIOS DO SERPRO": ("CNPJ", "🪪", "Comprovantes de inscrição no CNPJ da sede, regionais e escritórios."),
    "DECLARACOES": ("Declarações", "✍️", "Declarações exigidas em contratações (menores, etc.)."),
    "INFORMACOES CADASTRAIS DO SERPRO": ("Dados cadastrais", "🗂️", "Informações cadastrais do Serpro."),
    "VENCIMENTO DAS CERTIDOES NEGATIVAS": ("Controle de vencimentos", "📅", "Tabela de vencimento das certidões negativas."),
}

UFS = {  # sigla: (nome, capital)
    "AC": ("Acre", "Rio Branco"), "AL": ("Alagoas", "Maceió"), "AP": ("Amapá", "Macapá"), "AM": ("Amazonas", "Manaus"),
    "BA": ("Bahia", "Salvador"), "CE": ("Ceará", "Fortaleza"), "DF": ("Distrito Federal", "Brasília"), "ES": ("Espírito Santo", "Vitória"),
    "GO": ("Goiás", "Goiânia"), "MA": ("Maranhão", "São Luís"), "MT": ("Mato Grosso", "Cuiabá"), "MS": ("Mato Grosso do Sul", "Campo Grande"),
    "MG": ("Minas Gerais", "Belo Horizonte"), "PA": ("Pará", "Belém"), "PB": ("Paraíba", "João Pessoa"), "PR": ("Paraná", "Curitiba"),
    "PE": ("Pernambuco", "Recife"), "PI": ("Piauí", "Teresina"), "RJ": ("Rio de Janeiro", "Rio de Janeiro"), "RN": ("Rio Grande do Norte", "Natal"),
    "RS": ("Rio Grande do Sul", "Porto Alegre"), "RO": ("Rondônia", "Porto Velho"), "RR": ("Roraima", "Boa Vista"), "SC": ("Santa Catarina", "Florianópolis"),
    "SP": ("São Paulo", "São Paulo"), "SE": ("Sergipe", "Aracaju"), "TO": ("Tocantins", "Palmas"),
}


def categoria_de(caminho):
    topo = caminho.strip("/").split("/")[0]
    norm = lambda x: re.sub(r"[^A-Z0-9]+", " ", sem_acento(x).upper()).strip()
    limpo = norm(re.sub(r"^\d+\s*-\s*", "", topo))
    for k, v in CATEGORIAS.items():
        if norm(k) in limpo:
            return k, v
    return limpo, (topo.title(), "📄", "")


def local_de(caminho, cat):
    """UF e cidade a partir das pastas e do nome do arquivo (inclui grafias da pasta, ex. 'MATO CROSSO DO SUL')."""
    alvo = sem_acento(caminho).upper().replace("CROSSO", "GROSSO")
    m = re.search(r"[- ]([A-Z]{2})(?:\.PDF)?$", alvo.rsplit("/", 1)[-1].rsplit(".", 1)[0] + ".PDF")
    if m and m.group(1) in UFS:
        uf = m.group(1)
        return uf, UFS[uf][1]
    # nomes de estado (os mais longos primeiro: "MATO GROSSO DO SUL" antes de "MATO GROSSO")
    for uf, (nome, capital) in sorted(UFS.items(), key=lambda x: -len(x[1][0])):
        if cat == "CERTIDAO DE REGULARIDADE ESTADUAL" and sem_acento(nome).upper() in alvo:
            return uf, None
    for uf, (nome, capital) in sorted(UFS.items(), key=lambda x: -len(x[1][1])):
        if re.search(r"\b" + re.escape(sem_acento(capital).upper()) + r"\b", alvo):
            return uf, capital
    if "SEDE" in alvo:
        return "DF", "Brasília (sede)"
    return None, None


PALAVRAS = {
    "CERTIDAO": "Certidão", "CERTIDOES": "Certidões", "DECLARACAO": "Declaração", "ESCRITORIO": "Escritório", "INFORMACOES": "Informações",
    "FALENCIA": "Falência", "FALENCIAS": "Falências", "RECUPERACOES": "Recuperações", "RECUPERACAO": "Recuperação", "DEBITOS": "Débitos",
    "CONTRATACAO": "Contratação", "BALANCO": "Balanço", "BALANCOS": "Balanços", "SAO": "São", "BELEM": "Belém", "BRASILIA": "Brasília",
    "GOIANIA": "Goiânia", "VITORIA": "Vitória", "MACEIO": "Maceió", "MACAPA": "Macapá", "CUIABA": "Cuiabá", "JOAO": "João", "LUIS": "Luís",
    "FLORIANOPOLIS": "Florianópolis", "PARAIBA": "Paraíba", "PARANA": "Paraná", "PIAUI": "Piauí", "CEARA": "Ceará", "ESPIRITO": "Espírito",
    "AMAPA": "Amapá", "RONDONIA": "Rondônia", "MARANHAO": "Maranhão", "GOIAS": "Goiás", "PERNAMPUCO": "Pernambuco", "DIVIDA": "Dívida",
    "INIDONEO": "Inidôneos", "ESPECIFICA": "Específica", "JURIDICA": "Jurídica", "PATRIMONIAL": "Patrimonial",
}
SIGLAS = {"SICAF", "CNPJ", "TCU", "FGTS", "CNDT", "DRE", "PGFN", "SRFB", "INSS", "CADIN", "TJDFT", "TST", "PGE", "GDF", "CND", "RFB", "SERPRO", "IV", "II", "III"} | set(UFS)
MINUSCULAS = {"DE", "DA", "DO", "DAS", "DOS", "E", "EM", "PARA"}


def titulo(nome):
    base = nome.rsplit(".", 1)[0].replace("_", " ")
    palavras = []
    for i, w in enumerate(re.split(r"(\s+|-)", base)):
        chave_w = sem_acento(w).upper()
        if not w.strip() or w == "-":
            palavras.append(w)
        elif chave_w in PALAVRAS:
            palavras.append(PALAVRAS[chave_w])
        elif w.upper() in SIGLAS or chave_w in SIGLAS:
            palavras.append(w.upper())
        elif chave_w in MINUSCULAS and i > 0:
            palavras.append(w.lower())
        elif w.isupper() or w.islower():
            palavras.append(w[:1].upper() + w[1:].lower())
        else:
            palavras.append(w)
    t = re.sub(r"\s+", " ", "".join(palavras)).strip()
    t = re.sub(r"\b(do|de) para\b", r"\1 Pará", t)  # o estado, não a preposição
    return t.replace("Serpro(23)", "Serpro").replace(" (1)", "")


# ---------------------------------------------------------------- principal
def main():
    anterior = {}
    if DADOS.exists():
        for d in json.loads(DADOS.read_text("utf-8")).get("documentos", []):
            anterior[d["caminho"]] = d
    print("listando a pasta pública…")
    itens = arvore()
    arquivos = [i for i in itens if not i["pasta"]]
    print(f"{len(arquivos)} arquivos em {sum(i['pasta'] for i in itens)} pastas")

    # planilha de vencimentos primeiro (é fonte de validade para os demais)
    vencimentos = {}
    for i in arquivos:
        if i["caminho"].lower().endswith(".xlsx") and "VENCIMENTO" in sem_acento(i["caminho"]).upper():
            destino = CACHE / i["etag"] / "tabela.xlsx"
            if not destino.exists():
                baixar(i["caminho"], destino)
            vencimentos = tabela_vencimentos(destino)

    documentos, novos, reaproveitados = [], 0, 0
    for i in arquivos:
        caminho = i["caminho"]
        nome = caminho.rsplit("/", 1)[1]
        ext = nome.rsplit(".", 1)[-1].lower() if "." in nome else ""
        if nome.lower() == "readme.md" and i["tamanho"] < 5:
            continue  # Leia-me vazio
        id_ = hashlib.sha1(caminho.encode()).hexdigest()[:12]
        pasta = caminho.rsplit("/", 1)[0] or "/"
        cat_id, (cat_nome, _icone, _desc) = categoria_de(caminho)
        uf, cidade = local_de(caminho, cat_id)
        velho = anterior.get(caminho)
        doc = {
            "id": id_, "caminho": caminho, "arquivo": nome, "titulo": titulo(nome), "tipo": ext,
            "pasta": pasta, "pastas": [p for p in pasta.strip("/").split("/") if p],
            "categoria": cat_id, "uf": uf, "cidade": cidade,
            "tamanho": i["tamanho"], "modificadoEm": i["modificadoEm"], "etag": i["etag"],
            "primeiroVisto": (velho or {}).get("primeiroVisto") or dt.datetime.utcnow().replace(microsecond=0).isoformat() + "Z",
            "download": link_download(caminho), "abrirPasta": link_pasta(pasta),
        }
        pdf_local = SITE / "arquivos" / f"{id_}.pdf"
        mini = SITE / "miniaturas" / f"{id_}.webp"
        mesmo = velho and velho.get("etag") == i["etag"] and (ext != "pdf" or pdf_local.exists())
        if mesmo and velho.get("leitura") == "sem texto" and shutil.which("tesseract"):
            mesmo = False  # escaneado lido antes sem OCR: agora há leitura de imagem, lê de novo
        if mesmo:
            for k in ("texto", "leitura", "validadeDoc", "urlExterna", "miniatura", "ver"):
                if k in velho:
                    doc[k] = velho[k]
            reaproveitados += 1
        else:
            novos += 1
            print("   lendo", caminho)
            local = CACHE / i["etag"] / nome
            if not local.exists():
                baixar(caminho, local)
            if ext == "pdf":
                texto, leitura = texto_pdf(local)
                doc["texto"] = re.sub(r"[ \t]+", " ", texto).strip()[:8000]
                doc["leitura"] = leitura
                doc["validadeDoc"] = validade_no_texto(texto, cat_id)
                pdf_local.parent.mkdir(parents=True, exist_ok=True)
                shutil.copyfile(local, pdf_local)
                doc["ver"] = f"arquivos/{id_}.pdf"
                if miniatura(local, mini):
                    doc["miniatura"] = f"miniaturas/{id_}.webp"
            elif ext == "url":
                m = re.search(r"^URL=(.+)$", local.read_text("latin-1", "ignore"), re.M)
                doc["urlExterna"] = m.group(1).strip() if m else ""
                doc["texto"] = doc["urlExterna"]
            elif ext == "md":
                doc["texto"] = local.read_text("utf-8", "ignore")[:4000]
            elif ext == "xlsx":
                doc["texto"] = "Tabela de controle de vencimento das certidões negativas do Serpro (municipais e estaduais)."

        # validade: o documento manda; a planilha completa quando o PDF não traz a data
        validade = doc.get("validadeDoc")
        if not validade and cat_id in ("CERTIDAO DE REGULARIDADE MUNICIPAL", "CERTIDAO DE REGULARIDADE ESTADUAL") and vencimentos:
            secao = "municipal" if "MUNICIPAL" in cat_id else "estadual"
            alvo = chave(cidade or "") if secao == "municipal" else chave(UFS[uf][0] if uf else "")
            if "SEDE" in sem_acento(caminho).upper() and secao == "municipal":
                alvo = "SERPROSEDE"
            ate = vencimentos.get((secao, alvo))
            if ate:
                validade = {"ate": ate, "fonte": "tabela de vencimentos"}
        if validade:
            doc["validade"] = validade
        documentos.append(doc)

    # limpa cópias de arquivos que saíram da pasta
    ids = {d["id"] for d in documentos}
    for sub in ("arquivos", "miniaturas"):
        for f in (SITE / sub).glob("*"):
            if f.stem not in ids:
                f.unlink()

    contagem = {}
    for d in documentos:
        contagem[d["categoria"]] = contagem.get(d["categoria"], 0) + 1
    categorias = []
    for d in documentos:
        if d["categoria"] in [c["id"] for c in categorias]:
            continue
        _, (nome, icone, desc) = categoria_de(d["caminho"])
        categorias.append({"id": d["categoria"], "nome": nome, "icone": icone, "descricao": desc, "total": contagem[d["categoria"]],
                           "ordem": d["caminho"].strip("/").split("/")[0], "abrirPasta": link_pasta("/" + d["caminho"].strip("/").split("/")[0])})
    categorias.sort(key=lambda c: c["ordem"])

    documentos.sort(key=lambda d: d["caminho"])
    saida = {
        "geradoEm": dt.datetime.utcnow().replace(microsecond=0).isoformat() + "Z",
        "fonte": {"nome": "SUPCO_TRIBUTARIO_DOCUMENTAÇÃO", "url": f"{BASE}/s/{TOKEN}", "zip": f"{BASE}/s/{TOKEN}/download"},
        "categorias": categorias,
        "documentos": documentos,
    }
    novo_json = json.dumps(saida, ensure_ascii=False, indent=0, sort_keys=True)  # ordem fixa: só muda o que mudou de fato
    # só regrava se algo além do horário mudou (evita commit a cada hora sem novidade)
    if DADOS.exists():
        velho = json.loads(DADOS.read_text("utf-8")); velho["geradoEm"] = saida["geradoEm"]
        if json.dumps(velho, ensure_ascii=False, indent=0, sort_keys=True) == novo_json:
            print("nada mudou na pasta")
            return
    DADOS.write_text(novo_json, "utf-8")
    com_validade = sum(1 for d in documentos if d.get("validade"))
    print(f"ok: {len(documentos)} documentos ({novos} lidos, {reaproveitados} reaproveitados), {com_validade} com validade")


if __name__ == "__main__":
    main()
